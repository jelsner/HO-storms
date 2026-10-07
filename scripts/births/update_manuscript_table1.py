from __future__ import annotations

import csv
import re
import shutil
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "manuscript" / "births" / "Final Manuscript Draft 4.1.26.docx"
VALUES = ROOT / "outputs" / "births" / "revised_table1" / "revised_table1_values.csv"
OUTPUT = (
    ROOT
    / "outputs"
    / "births"
    / "revised_table1"
    / "Final Manuscript Draft 4.1.26 - Revised Table 1.docx"
)


def normalize_label(value: str) -> str:
    value = value.replace("’", "'").replace("“", '"').replace("”", '"')
    return re.sub(r"\s+", " ", value).strip()


def set_cell_text(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    paragraph_properties = deepcopy(paragraph._p.pPr) if paragraph._p.pPr is not None else None
    run_properties = None
    for existing_paragraph in cell.paragraphs:
        for run in existing_paragraph.runs:
            if run._r.rPr is not None:
                run_properties = deepcopy(run._r.rPr)
                break
        if run_properties is not None:
            break

    for child in list(cell._tc):
        if child.tag == qn("w:p"):
            cell._tc.remove(child)

    new_paragraph = OxmlElement("w:p")
    if paragraph_properties is not None:
        new_paragraph.append(paragraph_properties)
    cell._tc.append(new_paragraph)
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    if run_properties is not None:
        run._r.insert(0, run_properties)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    table_header = tr_pr.find(qn("w:tblHeader"))
    if table_header is None:
        table_header = OxmlElement("w:tblHeader")
        table_header.set(qn("w:val"), "true")
        tr_pr.append(table_header)


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE, OUTPUT)

    with VALUES.open(newline="", encoding="utf-8") as handle:
        values = list(csv.DictReader(handle))
    if len(values) != 79:
        raise RuntimeError(f"Expected 79 Table 1 rows, found {len(values)}")

    document = Document(OUTPUT)
    table = document.tables[0]
    if len(table.rows) != 83 or len(table.columns) != 5:
        raise RuntimeError(
            f"Unexpected source Table 1 shape: {len(table.rows)} rows x {len(table.columns)} columns"
        )

    set_cell_text(
        table.rows[0].cells[0],
        "Table 1: Study Population Characteristics by Corrected Tropical Cyclone Exposure",
    )
    set_cell_text(table.rows[1].cells[1], "Overall")
    set_cell_text(table.rows[1].cells[2], "Any Storm Effect")
    set_cell_text(table.rows[1].cells[3], "No Storm Effect")
    set_cell_text(table.rows[1].cells[4], "p-value")
    set_cell_text(table.rows[2].cells[0], "Variables")
    set_cell_text(table.rows[2].cells[1], "N = 4,484,317")
    set_cell_text(table.rows[2].cells[2], "N = 1,685,418")
    set_cell_text(table.rows[2].cells[3], "N = 2,798,899")
    set_cell_text(table.rows[2].cells[4], "")
    for row_index, value_row in enumerate(values, start=3):
        source_label = normalize_label(table.rows[row_index].cells[0].text)
        calculated_label = normalize_label(value_row["label"])
        if source_label != calculated_label:
            raise RuntimeError(
                f"Row {row_index}: source label {source_label!r} does not match calculated label {calculated_label!r}"
            )
        set_cell_text(table.rows[row_index].cells[1], value_row["overall"])
        set_cell_text(table.rows[row_index].cells[2], value_row["exposed"])
        set_cell_text(table.rows[row_index].cells[3], value_row["unexposed"])
        set_cell_text(table.rows[row_index].cells[4], value_row["p_value"])

    note_cell = table.rows[82].cells[0]
    for other_cell in table.rows[82].cells[1:]:
        note_cell = note_cell.merge(other_cell)
    note = (
        "Note: Values are mean (SD) or n (%). Any Storm Effect denotes at least one corrected "
        "residence-specific 34-kt +/-7-day tropical cyclone exposure window overlapping the first, "
        "second, or third trimester. The cohort follows the revised model eligibility criteria. "
        "P-values use Welch's t-test for maternal age and Pearson chi-square tests for categorical variables."
    )
    set_cell_text(table.rows[82].cells[0], note)
    note_paragraph = table.rows[82].cells[0].paragraphs[0]
    note_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    note_paragraph.paragraph_format.space_before = Pt(3)
    note_paragraph.paragraph_format.space_after = Pt(3)
    for run in note_paragraph.runs:
        run.italic = True
        run.font.size = Pt(8)

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
